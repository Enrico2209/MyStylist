#!/usr/bin/env python3
"""
Rigenera il pool di outfit tenendo conto delle correzioni fatte a mano in
nuvolari_db/capi_parametri.xlsx (vedi esporta_capi_excel.py).

Una cella conta come modificata se differisce dal valore salvato nel foglio
nascosto "originali" al momento dell'esportazione. Solo quelle celle vengono
applicate: tutto il resto resta com'e' nella pipeline, anche se nel frattempo
e' stata rifatta.

Le modifiche si applicano in due tempi, perche' la pipeline ricava i
parametri in due posti diversi:
  - stile, formalita', stagione, fantasia, colore, abbinabile stanno nel
    parquet delle feature -> si correggono li', PRIMA del clustering (stile e
    formalita' decidono il cluster);
  - slot, genere, manica, gamba e le quattro regole di forma li deduce
    load_and_prepare dal titolo -> si sovrascrivono DOPO.

features_clustered.parquet non viene toccato: la versione corretta va in
features_clustered_modificate.parquet. Il pool prima di ogni run viene
copiato in outfits_pool.jsonl.prima_excel_<data>.bak.

Uso:
    python3 genera_da_excel.py --controlla   # elenca le modifiche, non genera
    python3 genera_da_excel.py               # applica e rigenera il pool (ore di CPU)
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from skimage.color import rgb2lab

sys.path.insert(0, str(Path(__file__).resolve().parent))
from percorsi import DATI  # noqa: E402
from esporta_capi_excel import (  # noqa: E402
    COLONNE, MODIFICABILI, SCELTE, EXCEL, CLUSTERED)
from feature_engineering import STYLE_TAGS  # noqa: E402

MODIFICATE = DATI / "features_modificate.parquet"
CLUSTERED_MODIFICATE = DATI / "features_clustered_modificate.parquet"
POOL = DATI / "outfits_pool.jsonl"
IMMAGINI = DATI / "outfit_images"

# Dove agisce ciascuna colonna: nel parquet (prima del clustering) o sul df
# di load_and_prepare (dopo).
PRIMA = {"abbinabile", "stagione", "formalita", "fantasia", "colore",
         *[f"stile_{t}" for t in STYLE_TAGS]}
DOPO = {"slot": "slot", "genere": "gender", "manica": "sleeve", "gamba": "leg_length",
        "mocassino": "mocassino", "top_senza_collo": "top_senza_collo",
        "cappotto": "cappotto", "bottom_da_tuta": "bottom_da_tuta"}
NOME = {k: intest for intest, k, _, _ in COLONNE}


def normalizza(chiave, valore):
    """Valore della cella nella forma canonica; ValueError se non e' ammesso."""
    if valore is None or (isinstance(valore, str) and not valore.strip()):
        raise ValueError("cella vuota")
    if chiave in SCELTE:
        v = str(valore).strip().lower()
        if "sì" in SCELTE[chiave]:
            v = {"si": "sì", "s": "sì", "true": "sì", "vero": "sì",
                 "false": "no", "falso": "no", "n": "no"}.get(v, v)
        if v not in SCELTE[chiave]:
            raise ValueError(f"'{valore}' non e' fra {', '.join(SCELTE[chiave])}")
        return v
    if chiave == "formalita":
        f = float(str(valore).replace(",", "."))
        if f != int(f) or not 1 <= f <= 5:
            raise ValueError(f"'{valore}' non e' un intero da 1 a 5")
        return int(f)
    if chiave.startswith("stile_"):
        s = float(str(valore).replace(",", "."))
        if not 0 <= s <= 1:
            raise ValueError(f"'{valore}' non e' fra 0 e 1")
        return round(s, 2)
    if chiave == "colore":
        h = str(valore).strip().upper().lstrip("#")
        if len(h) != 6 or any(c not in "0123456789ABCDEF" for c in h):
            raise ValueError(f"'{valore}' non e' un colore esadecimale tipo #1F2A44")
        return "#" + h
    raise KeyError(chiave)


def leggi_excel():
    """{relpath: {chiave: (vecchio, nuovo)}} delle sole celle cambiate, ed errori."""
    wb = load_workbook(EXCEL, data_only=True)
    if "capi" not in wb.sheetnames or "originali" not in wb.sheetnames:
        sys.exit(f"[ERRORE] {EXCEL.name}: mancano i fogli 'capi' o 'originali'. Riesportalo.")

    orig = wb["originali"].iter_rows(values_only=True)
    intest_orig = next(orig)
    originali = {r[0]: dict(zip(intest_orig[1:], r[1:])) for r in orig if r[0]}

    ws = wb["capi"]
    intest = [c.value for c in ws[1]]
    colonna = {}
    for nome, chiave, _, _ in COLONNE:
        if nome not in intest:
            sys.exit(f"[ERRORE] Nel foglio 'capi' manca la colonna '{nome}'.")
        colonna[chiave] = intest.index(nome)

    modifiche, errori, sconosciuti = {}, [], []
    for n, riga in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        relpath = riga[colonna["relpath"]]
        if not relpath:
            continue
        if relpath not in originali:
            sconosciuti.append(relpath)
            continue
        for k in MODIFICABILI:
            try:
                nuovo = normalizza(k, riga[colonna[k]])
            except ValueError as e:
                errori.append(f"riga {n}, {NOME[k]}: {e}")
                continue
            vecchio = normalizza(k, originali[relpath][k])
            if nuovo != vecchio:
                modifiche.setdefault(relpath, {})[k] = (vecchio, nuovo)
    return modifiche, errori, sconosciuti


def hex_a_lab(h):
    rgb = np.array([[[int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)]]])
    return [float(x) for x in rgb2lab(rgb)[0, 0]]


def applica_al_parquet(feat: pd.DataFrame, modifiche: dict) -> pd.DataFrame:
    feat = feat.copy()
    idx = {r: i for i, r in enumerate(feat["relpath"])}
    for relpath, mod in modifiche.items():
        i = idx[relpath]
        if "abbinabile" in mod:
            feat.at[i, "needs_vision_review"] = mod["abbinabile"][1] == "no"
        if "stagione" in mod:
            for s in ("estate", "inverno", "mezza_stagione", "tutte"):
                feat.at[i, f"season_{s}"] = 1.0 if s == mod["stagione"][1] else 0.0
        if "formalita" in mod:
            feat.at[i, "formality_norm"] = (mod["formalita"][1] - 1) / 4
        if "fantasia" in mod:
            feat.at[i, "pattern"] = mod["fantasia"][1]
        if "colore" in mod:
            L, a, b = hex_a_lab(mod["colore"][1])
            feat.at[i, "L"], feat.at[i, "a"], feat.at[i, "b"] = L, a, b
            feat.at[i, "color_palette"] = json.dumps([{"lab": [L, a, b], "proportion": 1.0}])
        stili = [k for k in mod if k.startswith("stile_")]
        if stili:
            # Si prendono tutti e cinque, cambiati o no, e si riportano a somma 1:
            # e' la forma in cui feature_engineering li produce.
            valori = {t: mod.get(f"stile_{t}", (None, None))[1] for t in STYLE_TAGS}
            for t in STYLE_TAGS:
                if valori[t] is None:
                    valori[t] = float(feat.at[i, f"style_{t}"])
            totale = sum(valori.values())
            if totale <= 0:
                raise ValueError(f"{relpath}: gli stili sono tutti a zero")
            for t in STYLE_TAGS:
                feat.at[i, f"style_{t}"] = valori[t] / totale
    return feat


def applica_dopo(df: pd.DataFrame, modifiche: dict) -> pd.DataFrame:
    idx = {r: i for i, r in zip(df.index, df["relpath"])}
    for relpath, mod in modifiche.items():
        i = idx[relpath]
        for k, col in DOPO.items():
            if k not in mod:
                continue
            v = mod[k][1]
            if k == "slot":
                v = None if v == "nessuno" else v
            elif k == "genere":
                v = None if v == "unisex" else v
            elif k in ("mocassino", "top_senza_collo", "cappotto", "bottom_da_tuta"):
                v = v == "sì"
            df.at[i, col] = v
    return df


def id_outfit(percorso: Path) -> set:
    if not percorso.exists():
        return set()
    return {json.loads(l)["outfit_id"] for l in open(percorso, encoding="utf-8") if l.strip()}


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--controlla", action="store_true",
                    help="elenca le modifiche trovate e si ferma, senza generare")
    args = ap.parse_args()

    if not EXCEL.exists():
        sys.exit(f"[ERRORE] {EXCEL} non esiste: lancia prima esporta_capi_excel.py")

    modifiche, errori, sconosciuti = leggi_excel()
    feat = pd.read_parquet(CLUSTERED)
    fuori = [r for r in modifiche if r not in set(feat["relpath"])]

    print(f"[*] {EXCEL.name}: {sum(len(m) for m in modifiche.values())} celle modificate "
          f"su {len(modifiche)} capi")
    for relpath, mod in sorted(modifiche.items()):
        print(f"  {relpath}")
        for k, (vecchio, nuovo) in mod.items():
            print(f"      {NOME[k]}: {vecchio} -> {nuovo}")
    if sconosciuti:
        print(f"[!] {len(sconosciuti)} righe con una 'Capo (chiave)' che non c'era all'esportazione: ignorate")
    if fuori:
        print(f"[!] {len(fuori)} capi modificati non sono piu' nella pipeline: ignorati")
        for r in fuori:
            modifiche.pop(r)
    if errori:
        print(f"\n[ERRORE] {len(errori)} celle con valori non ammessi, correggile e rilancia:")
        for e in errori:
            print("   ", e)
        sys.exit(1)
    if args.controlla:
        print("\n[controlla] niente generato. Per rigenerare: python3 genera_da_excel.py")
        return
    if not modifiche:
        print("[*] nessuna modifica: il pool si rigenera comunque con i parametri attuali.")

    inizio = datetime.now()
    prima = id_outfit(POOL)
    con_foto = {p.stem for p in IMMAGINI.glob("*.png")}
    if POOL.exists():
        copia = POOL.with_suffix(POOL.suffix + f".prima_excel_{inizio:%d%b_%H%M}".lower() + ".bak")
        copia.write_bytes(POOL.read_bytes())
        print(f"[*] copia di sicurezza del pool: {copia.name}")

    print("\n── 1/3  modifiche nel parquet + clustering ──────────────────────")
    applica_al_parquet(feat, modifiche).to_parquet(MODIFICATE, index=False)
    from clustering import run_clustering
    run_clustering(features_parquet=str(MODIFICATE), out_parquet=str(CLUSTERED_MODIFICATE))

    print("\n── 2/3  attributi dedotti + modifiche ───────────────────────────")
    import outfit_generation as og
    df = applica_dopo(og.load_and_prepare(str(CLUSTERED_MODIFICATE)), modifiche)

    print("\n── 3/3  pool di outfit ──────────────────────────────────────────")
    og.run_outfit_pipeline(df, out_jsonl=str(POOL))

    dopo = id_outfit(POOL)
    print("\n══════════════════════════════════════════════════════════════════")
    print(f"outfit: {len(prima)} -> {len(dopo)}")
    print(f"immagini: {len(con_foto)} pagate, {len(con_foto & dopo)} ancora valide, "
          f"{len(con_foto - dopo)} orfane (non toccate)")
    print(f"durata: {datetime.now() - inizio}")


if __name__ == "__main__":
    main()
