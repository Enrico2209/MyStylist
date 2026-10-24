#!/usr/bin/env python3
"""
Esporta in Excel i parametri di ogni capo, cosi' come li vede l'algoritmo
che compone gli outfit, perche' si possano correggere a mano.

Una riga per capo, con la foto. Le colonne a sfondo giallo sono quelle che
l'algoritmo usa e che si possono modificare; quelle grigie sono solo per
riconoscere il capo. Dopo le modifiche, genera_da_excel.py rilegge il file e
rigenera il pool di outfit tenendone conto.

Il foglio nascosto "originali" conserva i valori di partenza: e' cosi' che
genera_da_excel.py distingue una modifica fatta a mano da un valore che e'
cambiato perche' nel frattempo la pipeline e' stata rifatta. Non va toccato.

Uso:
    python3 esporta_capi_excel.py
    -> nuvolari_db/capi_parametri.xlsx
"""

import io
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image
from skimage.color import lab2rgb

sys.path.insert(0, str(Path(__file__).resolve().parent))
from percorsi import DATI, CATALOGO  # noqa: E402
import outfit_generation as og  # noqa: E402
from scoring import FANTASIA_FORZA  # noqa: E402
from feature_engineering import STYLE_TAGS  # noqa: E402

CLUSTERED = DATI / "features_clustered.parquet"
EXCEL = DATI / "capi_parametri.xlsx"

LATO_FOTO = 110          # px, lato lungo della miniatura nel foglio
ALTEZZA_RIGA = 86        # punti: la miniatura ci sta dentro con un po' di margine

# Valori ammessi per le colonne a scelta. "nessuno"/"unisex" stanno per il None
# del codice: slot non riconosciuto (il capo non entra in nessun outfit) e
# genere non determinato (il capo e' un jolly, compatibile con uomo e donna).
SLOT = ["top", "bottom", "shoes", "outerwear", "accessory", "dress_or_suit", "nessuno"]
GENERE = ["uomo", "donna", "unisex"]
STAGIONE = list(og.SEASONS)
LUNGHEZZA = ["corta", "lunga"]
FANTASIA = list(FANTASIA_FORZA)
SI_NO = ["sì", "no"]

# (intestazione, chiave interna, modificabile, larghezza colonna)
COLONNE = [
    ("Foto",              "foto",            False, 16),
    ("Capo (chiave)",     "relpath",         False, 34),
    ("Titolo",            "title",           False, 38),
    ("Pagina negozio",    "url",             False, 18),
    ("Cluster",           "style_cluster",   False, 9),
    ("Abbinabile",        "abbinabile",      True, 11),
    ("Slot",              "slot",            True, 13),
    ("Genere",            "genere",          True, 10),
    ("Stagione",          "stagione",        True, 15),
    ("Manica",            "manica",          True, 9),
    ("Gamba",             "gamba",           True, 9),
    ("Formalità (1-5)",   "formalita",       True, 11),
    *[(f"Stile: {t}", f"stile_{t}", True, 11) for t in STYLE_TAGS],
    ("Fantasia",          "fantasia",        True, 13),
    ("Colore (hex)",      "colore",          True, 12),
    ("Mocassino",         "mocassino",       True, 11),
    ("Top senza collo",   "top_senza_collo", True, 11),
    ("Cappotto",          "cappotto",        True, 10),
    ("Pantalone da tuta", "bottom_da_tuta",  True, 11),
]
MODIFICABILI = [k for _, k, m, _ in COLONNE if m]

SCELTE = {
    "abbinabile": SI_NO, "slot": SLOT, "genere": GENERE, "stagione": STAGIONE,
    "manica": LUNGHEZZA, "gamba": LUNGHEZZA, "fantasia": FANTASIA,
    "mocassino": SI_NO, "top_senza_collo": SI_NO, "cappotto": SI_NO, "bottom_da_tuta": SI_NO,
}


def lab_a_hex(L, a, b) -> str:
    rgb = lab2rgb(np.array([[[L, a, b]]], dtype=float))[0, 0]
    return "#" + "".join(f"{int(round(max(0.0, min(1.0, c)) * 255)):02X}" for c in rgb)


def si_no(v) -> str:
    return "sì" if bool(v) else "no"


def valori_del_capo(r) -> dict:
    """I parametri del capo nella forma in cui compaiono nel foglio."""
    v = {
        "abbinabile": si_no(not r["needs_vision_review"]),
        "slot": r["slot"] or "nessuno",
        "genere": r["gender"] or "unisex",
        "stagione": r["season"],
        "manica": r["sleeve"],
        "gamba": r["leg_length"],
        "formalita": int(round(r["formality_norm"] * 4 + 1)),
        "fantasia": r["pattern"] or "tinta_unita",
        "colore": lab_a_hex(r["L"], r["a"], r["b"]),
        "mocassino": si_no(r["mocassino"]),
        "top_senza_collo": si_no(r["top_senza_collo"]),
        "cappotto": si_no(r["cappotto"]),
        "bottom_da_tuta": si_no(r["bottom_da_tuta"]),
    }
    for t in STYLE_TAGS:
        v[f"stile_{t}"] = round(float(r[f"style_{t}"]), 2)
    return v


def miniatura(r):
    """JPEG piccolo della foto principale; se e' illeggibile, la successiva."""
    candidate = [r["display_image"], *list(r["all_images"] if r["all_images"] is not None else [])]
    for rel in candidate:
        if not rel:
            continue
        try:
            with Image.open(CATALOGO / rel) as im:
                im = im.convert("RGB")
                im.thumbnail((LATO_FOTO, LATO_FOTO), Image.LANCZOS)
                buf = io.BytesIO()
                im.save(buf, "JPEG", quality=80)
                return buf, im.size
        except OSError:
            continue
    return None, None


def foglio_leggimi(wb):
    ws = wb.create_sheet("LEGGIMI", 0)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    righe = [
        ("COME SI USA", ""),
        ("1.", "Modifica solo le celle delle colonne con l'intestazione GIALLA, nel foglio 'capi'."),
        ("2.", "Salva il file in formato .xlsx, con lo stesso nome e nella stessa cartella."),
        ("3.", "Dalla cartella nuvolari_generation_pipeline lancia:  python3 genera_da_excel.py --controlla"),
        ("", "Elenca le modifiche che ha trovato, senza generare niente: serve a verificare che le abbia lette."),
        ("4.", "Poi lancia:  python3 genera_da_excel.py"),
        ("", "Rifà il clustering e rigenera il pool di outfit (nuvolari_db/outfits_pool.jsonl). Ci vogliono ore di CPU,"),
        ("", "nessun costo. Prima fa una copia di sicurezza del pool. Le foto degli outfit NON vengono generate."),
        ("", ""),
        ("ATTENZIONE", "Non toccare il foglio nascosto 'originali' e non cambiare la colonna 'Capo (chiave)'."),
        ("", "Le modifiche valgono solo per genera_da_excel.py: rigenera_tutto.py, catena.sh e il notebook non le leggono."),
        ("", "Per ripartire da zero, riesporta il file con python3 esporta_capi_excel.py (le modifiche fatte si perdono)."),
        ("", ""),
        ("COLONNE", ""),
        ("Abbinabile", "no = il capo non entra in nessun outfit."),
        ("Slot", "Il ruolo nell'outfit: top, bottom, shoes (scarpe), outerwear (capospalla), accessory."),
        ("", "Top, bottom e scarpe sono obbligatori; capospalla e accessorio facoltativi. dress_or_suit e nessuno"),
        ("", "non entrano negli outfit."),
        ("Genere", "uomo / donna / unisex. Unisex = compatibile con entrambi."),
        ("Stagione", "estate / inverno / mezza_stagione / tutte. Due capi stanno insieme se la stagione è uguale"),
        ("", "o se uno dei due è 'tutte'. Un capo invernale non sta con maniche o gambe corte."),
        ("Manica", "corta / lunga. Conta solo per i top: con un bermuda il top deve avere manica corta."),
        ("Gamba", "corta / lunga. Conta solo per i bottom: niente capospalla sopra un pantalone corto."),
        ("Formalità (1-5)", "1 = sportivo/da spiaggia, 5 = cerimonia. In un outfit i capi possono differire di un solo livello."),
        ("Stile: ...", "Quanto il capo appartiene a ciascun registro, da 0 a 1. Si rinormalizzano a somma 1:"),
        ("", "0,5 / 0,5 e 1 / 1 sono la stessa cosa. Insieme alla formalità decidono il cluster e metà del punteggio."),
        ("Fantasia", "tinta_unita, righe, quadri, stampato, camouflage, animalier. Due fantasie forti nello stesso"),
        ("", "outfit abbassano il punteggio; con una tinta unita nessuna penalità."),
        ("Colore (hex)", "Il colore dominante, es. #1F2A44. La cella è colorata così. L'algoritmo usa la tavolozza"),
        ("", "completa (fino a 3 colori): se cambi l'hex, la tavolozza viene sostituita da quel solo colore."),
        ("Mocassino", "sì = niente t-shirt, canotte o felpe insieme a questa scarpa."),
        ("Top senza collo", "sì = t-shirt, canotta, felpa: non sta con i mocassini."),
        ("Cappotto", "sì = caban, cappotto, trench: non sta sopra un pantalone da tuta."),
        ("Pantalone da tuta", "sì = jogger, tuta: niente cappotto sopra."),
        ("Cluster", "Solo informativo: viene ricalcolato da stile e formalità quando rigeneri."),
        ("", "-1 = capo fuori dai gruppi, non entra come capospalla o accessorio."),
    ]
    for i, (a, b) in enumerate(righe, start=1):
        ws.cell(i, 1, a).font = Font(bold=bool(a))
        ws.cell(i, 2, b)


def main():
    df = og.load_and_prepare(str(CLUSTERED))
    ordine_slot = {s: i for i, s in enumerate(SLOT)}
    df["_ordine"] = df["slot"].fillna("nessuno").map(ordine_slot)
    df = df.sort_values(["_ordine", "relpath"]).reset_index(drop=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "capi"
    orig = wb.create_sheet("originali")

    giallo = PatternFill("solid", fgColor="FFF2B3")
    grigio = PatternFill("solid", fgColor="E3E3E3")
    for j, (intest, chiave, modificabile, larghezza) in enumerate(COLONNE, start=1):
        c = ws.cell(1, j, intest)
        c.font = Font(bold=True)
        c.fill = giallo if modificabile else grigio
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[c.column_letter].width = larghezza
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "D2"

    orig.append(["relpath", *MODIFICABILI])
    indice = {k: j for j, (_, k, _, _) in enumerate(COLONNE, start=1)}
    senza_foto = []

    for i, r in df.iterrows():
        riga = i + 2
        v = valori_del_capo(r)
        ws.cell(riga, indice["relpath"], r["relpath"])
        ws.cell(riga, indice["title"], r["title"])
        link = ws.cell(riga, indice["url"], "apri" if r["url"] else "")
        if r["url"]:
            link.hyperlink = r["url"]
            link.font = Font(color="0563C1", underline="single")
        ws.cell(riga, indice["style_cluster"], int(r["style_cluster"]))
        for k in MODIFICABILI:
            ws.cell(riga, indice[k], v[k])
        ws.cell(riga, indice["colore"]).fill = PatternFill("solid", fgColor=v["colore"][1:])
        for k in MODIFICABILI + ["relpath", "title", "url", "style_cluster"]:
            ws.cell(riga, indice[k]).alignment = Alignment(vertical="center", wrap_text=(k == "title"))
        orig.append([r["relpath"], *[v[k] for k in MODIFICABILI]])

        buf, dim = miniatura(r)
        if buf is None:
            senza_foto.append(r["relpath"])
        else:
            img = XLImage(buf)
            img.width, img.height = dim
            ws.add_image(img, f"A{riga}")
        ws.row_dimensions[riga].height = ALTEZZA_RIGA
        if riga % 500 == 0:
            print(f"  {riga - 1}/{len(df)} capi", flush=True)

    ultima = len(df) + 1
    for k, scelte in SCELTE.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(scelte) + '"', allow_blank=False,
                            showErrorMessage=True, errorTitle="Valore non ammesso",
                            error="Scegli uno dei valori dell'elenco (vedi foglio LEGGIMI).")
        col = ws.cell(1, indice[k]).column_letter
        dv.add(f"{col}2:{col}{ultima}")
        ws.add_data_validation(dv)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5",
                        showErrorMessage=True, error="La formalità va da 1 a 5.")
    col = ws.cell(1, indice["formalita"]).column_letter
    dv.add(f"{col}2:{col}{ultima}")
    ws.add_data_validation(dv)
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1",
                        showErrorMessage=True, error="I valori di stile vanno da 0 a 1.")
    for t in STYLE_TAGS:
        col = ws.cell(1, indice[f"stile_{t}"]).column_letter
        dv.add(f"{col}2:{col}{ultima}")
    ws.add_data_validation(dv)
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLONNE)).column_letter}{ultima}"

    orig.sheet_state = "hidden"
    foglio_leggimi(wb)
    wb.active = 1
    wb.save(EXCEL)
    print(f"[OK] {len(df)} capi in {EXCEL.relative_to(DATI.parent)} ({EXCEL.stat().st_size / 1e6:.1f} MB)")
    if senza_foto:
        print(f"     senza foto leggibile: {len(senza_foto)}")
        for p in senza_foto:
            print("      ", p)


if __name__ == "__main__":
    main()
